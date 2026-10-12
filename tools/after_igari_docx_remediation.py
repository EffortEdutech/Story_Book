from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document


def iter_paragraphs(doc):
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def set_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def replace_exact(doc, old: str, new: str, changes: list[str]) -> bool:
    for paragraph in iter_paragraphs(doc):
        if old in paragraph.text:
            set_paragraph_text(paragraph, paragraph.text.replace(old, new))
            changes.append(f"{old!r} -> {new!r}")
            return True
    return False


def remove_exact_paragraph(doc, text: str, changes: list[str]) -> bool:
    for paragraph in iter_paragraphs(doc):
        if paragraph.text.strip() == text:
            p = paragraph._element
            p.getparent().remove(p)
            changes.append(f"removed paragraph {text!r}")
            return True
    return False


def blank_exact_paragraph(doc, text: str, changes: list[str]) -> bool:
    for paragraph in iter_paragraphs(doc):
        if paragraph.text.strip() == text:
            set_paragraph_text(paragraph, "")
            changes.append(f"blanked paragraph {text!r}")
            return True
    return False


def search(base: Path) -> None:
    terms = [
        "Behind him, his cold coffee",
        "For now.",
        "The pilots knew.",
        "Somebody somewhere had told them where to go.",
        "For one brief moment Adam could never quite describe",
        "The aircraft began to move at 00:31.",
        "Four minutes had passed",
        "lowered the tray table",
        "The last thing he thought of",
        "Then, finally, he stopped checking the time.",
        "Twelve minutes, perhaps",
        "The minute display had not changed.",
        "Not disappearing for a month.",
        "The stars had shifted relative to the frame.",
        "Still 01:20.",
        "before the minute could change",
        "Even if there had been a change in engine tone",
        "Steady. Then not steady. Then steady again.",
        "from shoulder to hip",
        "window frame",
        "star",
        "world could disappear completely",
        "Adam recognised it this time",
        "Adam noticed it this time",
        "The justification settled comfortably",
        "The explanation pleased him more than it should have",
    ]
    result = {}
    for path in sorted(base.glob("AFTER_IGARI_Chapter_*.docx")):
        doc = Document(path)
        hits = []
        for i, p in enumerate(iter_paragraphs(doc), start=1):
            text = p.text
            for term in terms:
                if term in text:
                    hits.append({"paragraph": i, "term": term, "text": text})
        if hits:
            result[path.name] = hits
    print(json.dumps(result, ensure_ascii=False, indent=2))


def remediate(base: Path) -> None:
    edits: dict[str, list[str]] = {}

    def open_doc(name: str):
        path = base / name
        return path, Document(path), edits.setdefault(name, [])

    # Chapter 1
    path, doc, changes = open_doc("AFTER_IGARI_Chapter_01_The_Message.docx")
    remove_exact_paragraph(
        doc,
        "Behind him, his cold coffee remained beneath the seat until a cleaner noticed it several minutes later.",
        changes,
    )
    replace_exact(doc, "In Adam's pocket, Hana's message stayed unopened. For now.", "In Adam's pocket, Hana's message stayed unopened.", changes)
    blank_exact_paragraph(doc, "For now.", changes)
    doc.save(path)

    # Chapter 4
    path, doc, changes = open_doc("AFTER_IGARI_Chapter_04_Before_the_Door_Closes_REISSUED.docx")
    blank_exact_paragraph(doc, "The justification settled comfortably.", changes)
    doc.save(path)

    # Chapter 5
    path, doc, changes = open_doc("AFTER_IGARI_Chapter_05_Takeoff.docx")
    replace_exact(doc, "The aircraft began to move at 00:31.", "The aircraft began to move.", changes)
    replace_exact(doc, "Four minutes had passed", "Five minutes had passed", changes)
    replace_exact(doc, "For one brief moment Adam could never quite describe", "For one brief moment Adam could not have described", changes)
    remove_exact_paragraph(doc, "The pilots knew. Somebody somewhere had told them where to go. That was enough.", changes)
    blank_exact_paragraph(doc, "The pilots knew.", changes)
    blank_exact_paragraph(doc, "Somebody somewhere had told them where to go.", changes)
    blank_exact_paragraph(doc, "That was enough.", changes)
    blank_exact_paragraph(doc, "Adam recognised it this time.", changes)
    blank_exact_paragraph(doc, "The explanation pleased him more than it should have.", changes)
    doc.save(path)

    # Chapter 6
    path, doc, changes = open_doc("AFTER_IGARI_Chapter_06_Cruise.docx")
    replace_exact(doc, "lowered the tray table", "raised the tray table", changes)
    replace_exact(doc, "The last thing he thought of", "The thing he thought of as he drifted", changes)
    blank_exact_paragraph(doc, "For now.", changes)
    doc.save(path)

    # Chapter 7
    path, doc, changes = open_doc("AFTER_IGARI_Chapter_07_Thirty_Five_Thousand_Feet.docx")
    replace_exact(doc, "Then, finally, he stopped checking the time.", "Then, finally, he let his hand fall from the watch.", changes)
    replace_exact(doc, "his parents when they needed things done", "his mother when she needed things done", changes)
    doc.save(path)

    # Chapter 8
    path, doc, changes = open_doc("AFTER_IGARI_Chapter_08_Good_Night.docx")
    replace_exact(doc, "Twelve minutes, perhaps", "A few minutes, perhaps", changes)
    remove_exact_paragraph(doc, "The minute display had not changed.", changes)
    replace_exact(doc, "Not disappearing for a month.", "Not away for a month.", changes)
    replace_exact(doc, "The stars had shifted relative to the frame.", "The stars remained there, small and indifferent beyond the glass.", changes)
    replace_exact(doc, "before his parents needed help with appointments", "before his mother needed help with appointments", changes)
    replace_exact(doc, "where explanations went", "where the thought went", changes)
    doc.save(path)

    # Chapter 9
    path, doc, changes = open_doc("AFTER_IGARI_Chapter_09_IGARI.docx")
    remove_exact_paragraph(doc, "Still 01:20.", changes)
    replace_exact(doc, "lowered his hand before the minute could change", "lowered his hand", changes)
    replace_exact(doc, "Even if there had been a change in engine tone.", "Even if attention made repetition feel uneven.", changes)
    replace_exact(doc, "Even if there had been a change in engine tone. Had there?", "Even if attention made repetition feel uneven.", changes)
    replace_exact(doc, "Steady. Then not steady. Then steady again.", "Steady, because repetition could become uneven when a man listened too closely.", changes)
    replace_exact(doc, "from shoulder to hip", "through his body", changes)
    # Remove directional cues while preserving paragraph flow.
    for paragraph in iter_paragraphs(doc):
        text = paragraph.text
        new = re.sub(r"\s*The window frame[^.]*\.", "", text)
        new = re.sub(r"\s*Outside the window, one of the stars disappeared\.", "", new)
        new = new.replace("the world could disappear completely", "the world could fall silent")
        new = new.replace("Adam noticed it this time", "The word arrived again")
        if new != text:
            set_paragraph_text(paragraph, new)
            changes.append(f"adjusted Chapter 9 paragraph: {text!r}")
    blank_exact_paragraph(doc, "The window frame no longer looked level.", changes)
    blank_exact_paragraph(doc, "Outside the window, one of the stars disappeared.", changes)
    doc.save(path)

    print(json.dumps(edits, ensure_ascii=False, indent=2))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("mode", choices=["search", "remediate"])
    parser.add_argument("--base", required=True)
    args = parser.parse_args()
    base = Path(args.base)
    if args.mode == "search":
        search(base)
    else:
        remediate(base)


if __name__ == "__main__":
    main()
