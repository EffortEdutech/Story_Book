from __future__ import annotations

import json
from pathlib import Path

from docx import Document


BASE = Path("02_After_Igari/Manuscript")

BLOCKED = [
    "Behind him, his cold coffee remained beneath the seat",
    "For now.",
    "The aircraft began to move at 00:31.",
    "Four minutes had passed.",
    "For one brief moment Adam could never quite describe",
    "The pilots knew.",
    "Somebody somewhere had told them where to go.",
    "That was enough.",
    "The justification settled comfortably.",
    "Adam recognised it this time.",
    "The explanation pleased him more than it should have.",
    "lowered the tray table",
    "The last thing he thought of",
    "Then, finally, he stopped checking the time.",
    "Twelve minutes, perhaps.",
    "The minute display had not changed.",
    "Not disappearing for a month.",
    "The stars had shifted relative to the frame.",
    "Still 01:20.",
    "before the minute could change",
    "Even if there had been a change in engine tone",
    "Steady. Then not steady. Then steady again.",
    "from shoulder to hip",
    "The window frame no longer looked level.",
    "Outside the window, one of the stars disappeared.",
    "the world could disappear completely",
    "Adam noticed it this time.",
]


def iter_text(doc: Document):
    for p in doc.paragraphs:
        yield p.text
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p.text


def main() -> None:
    result = {}
    for path in sorted(BASE.glob("AFTER_IGARI_Chapter_*.docx")):
        doc = Document(path)
        text = "\n".join(iter_text(doc))
        hits = [term for term in BLOCKED if term in text]
        result[path.name] = {
            "paragraphs": len(list(iter_text(doc))),
            "blocked_hits": hits,
        }
    print(json.dumps(result, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
